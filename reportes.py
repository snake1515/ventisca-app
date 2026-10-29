"""Reportes consolidados por día, semana o mes, y exportación a Excel."""
import io
from collections import defaultdict
from datetime import date, datetime, timedelta

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from activos_calc import activos_vigentes, depreciacion_dia
from models import (TIPOS_SERVICIO, Activo, CompraInsumo, Equipo, Gasto, Orden, VentaEquipo, ahora)

MESES = ["ene", "feb", "mar", "abr", "may", "jun", "jul", "ago", "sep", "oct", "nov", "dic"]
AGRUPACIONES = {"diario": "Por día", "semanal": "Por semana", "mensual": "Por mes"}

COLUMNAS = [  # clave, título, es dinero
    ("num_ordenes", "Servicios hechos", False),
    ("servicios", "Ingresos servicios", True),
    ("costo_serv", "Materiales y técnicos", True),
    ("ventas_eq", "Venta de equipos", True),
    ("costo_eq", "Costo de equipos", True),
    ("bruta", "Ganancia bruta", True),
    ("equipos", "Depreciación herramientas", True),
    ("gastos", "Gastos", True),
    ("neta", "Ganancia neta", True),
    ("compras", "Compras materiales", True),
    ("compras_eq", "Compras herramientas", True),
]


def inicio_bucket(d, agrupar):
    if agrupar == "semanal":
        return d - timedelta(days=d.weekday())
    if agrupar == "mensual":
        return d.replace(day=1)
    return d


def etiqueta(d, agrupar):
    if agrupar == "semanal":
        fin = d + timedelta(days=6)
        sem = d.isocalendar()[1]
        if d.month == fin.month:
            return f"Sem {sem} · {d.day}–{fin.day} {MESES[d.month - 1]}"
        return f"Sem {sem} · {d.day} {MESES[d.month - 1]}–{fin.day} {MESES[fin.month - 1]}"
    if agrupar == "mensual":
        return f"{MESES[d.month - 1].capitalize()} {d.year}"
    return d.strftime("%d/%m/%Y")


def rango_por_defecto(agrupar, hoy):
    if agrupar == "mensual":
        m, y = hoy.month - 5, hoy.year
        while m < 1:
            m += 12
            y -= 1
        return date(y, m, 1), hoy
    if agrupar == "semanal":
        return inicio_bucket(hoy, "semanal") - timedelta(weeks=7), hoy
    return hoy.replace(day=1), hoy


def ordenes_periodo(ini, fin):
    return (Orden.query.filter(Orden.estado == "terminada", Orden.anulada.is_(False),
                               Orden.fecha_fin >= ini, Orden.fecha_fin < fin)
            .order_by(Orden.fecha_fin).all())


def ventas_eq_periodo(ini, fin):
    return (VentaEquipo.query.filter(VentaEquipo.anulada.is_(False), VentaEquipo.fecha >= ini,
                                     VentaEquipo.fecha < fin).order_by(VentaEquipo.fecha).all())


def generar(desde, hasta, agrupar):
    """desde/hasta: date (inclusive). Devuelve dict con filas, totales y detalles."""
    ini = datetime.combine(desde, datetime.min.time())
    fin = datetime.combine(hasta + timedelta(days=1), datetime.min.time())
    filas = defaultdict(lambda: defaultdict(float))

    def b(d):
        return inicio_bucket(d, agrupar)

    por_tipo = defaultdict(lambda: defaultdict(float))
    por_tecnico = defaultdict(lambda: defaultdict(float))
    por_cliente = defaultdict(lambda: defaultdict(float))
    ordenes = ordenes_periodo(ini, fin)
    for o in ordenes:
        f = filas[b(o.fecha_fin.date())]
        costo = o.costo_materiales + o.pago_tecnico + o.otros_costos
        f["num_ordenes"] += 1
        f["servicios"] += o.total
        f["costo_serv"] += costo
        t = por_tipo[o.tipo]
        t["n"] += 1
        t["equipos"] += len(o.equipos)
        t["ingresos"] += o.total
        t["ganancia"] += o.ganancia
        tec = por_tecnico[o.tecnico.nombre if o.tecnico else "Sin asignar"]
        tec["n"] += 1
        tec["ingresos"] += o.total
        tec["comision"] += o.pago_tecnico
        tec["ganancia"] += o.ganancia
        if o.en_garantia:
            tec["garantias"] += 1
        c = por_cliente[o.cliente_nombre or "—"]
        c["servicios"] += 1
        c["ingresos"] += o.total
    for v in ventas_eq_periodo(ini, fin):
        f = filas[b(v.fecha.date())]
        f["ventas_eq"] += v.precio_total
        f["costo_eq"] += v.costo_completo
        c = por_cliente[v.cliente_nombre or "—"]
        c["equipos"] += v.cantidad
        c["ingresos"] += v.precio_total
    for c in CompraInsumo.query.filter(CompraInsumo.anulada.is_(False), CompraInsumo.fecha >= ini,
                                       CompraInsumo.fecha < fin):
        filas[b(c.fecha.date())]["compras"] += c.costo_total
    for a in Activo.query.filter(Activo.eliminado.is_(False), Activo.fecha_compra >= desde,
                                 Activo.fecha_compra <= hasta):
        filas[b(a.fecha_compra)]["compras_eq"] += a.valor
    gastos_cat = defaultdict(float)
    for g in Gasto.query.filter(Gasto.fecha >= desde, Gasto.fecha <= hasta):
        filas[b(g.fecha)]["gastos"] += g.valor
        gastos_cat[g.categoria] += g.valor
    activos = activos_vigentes()
    d = desde
    while d <= hasta:
        dep = depreciacion_dia(d, activos)
        if dep:
            filas[b(d)]["equipos"] += dep
        d += timedelta(days=1)

    claves, k = [], b(desde)
    while k <= hasta:
        claves.append(k)
        if agrupar == "mensual":
            k = (k.replace(day=28) + timedelta(days=4)).replace(day=1)
        elif agrupar == "semanal":
            k += timedelta(weeks=1)
        else:
            k += timedelta(days=1)

    salida, tot = [], defaultdict(float)
    for k in claves:
        f = filas[k]
        f["ingresos"] = f["servicios"] + f["ventas_eq"]
        f["bruta"] = f["ingresos"] - f["costo_serv"] - f["costo_eq"]
        f["neta"] = f["bruta"] - f["equipos"] - f["gastos"]
        fila = {c: f[c] for c, _, _ in COLUMNAS}
        fila["ingresos"] = f["ingresos"]
        fila["margen"] = (f["neta"] / f["ingresos"] * 100) if f["ingresos"] else None
        fila["periodo"] = etiqueta(k, agrupar)
        fila["inicio"] = k
        salida.append(fila)
        for c in [c for c, _, _ in COLUMNAS] + ["ingresos"]:
            tot[c] += f[c]
    tot["margen"] = (tot["neta"] / tot["ingresos"] * 100) if tot["ingresos"] else None

    tipos = [dict(clave=k, nombre=TIPOS_SERVICIO.get(k, k), **v)
             for k, v in sorted(por_tipo.items(), key=lambda x: -x[1]["ingresos"])]
    tecnicos = [dict(nombre=k, **v) for k, v in sorted(por_tecnico.items(), key=lambda x: -x[1]["n"])]
    clientes = [dict(nombre=k, **v) for k, v in sorted(por_cliente.items(), key=lambda x: -x[1]["ingresos"])]
    return dict(filas=salida, tot=dict(tot), por_tipo=tipos, por_tecnico=tecnicos, por_cliente=clientes,
                gastos_cat=sorted(gastos_cat.items(), key=lambda x: -x[1]))


# ------------------------------------------------------------------ Excel
AZUL = "0D47A1"
CELESTE = "E3F2FD"
FMT_COP = '"$" #,##0;[Red]-"$" #,##0'
FMT_NUM = '#,##0'
FMT_PCT = '0.0"%"'
FMT_F = "DD/MM/YYYY"


def _encabezado(ws, fila, titulos):
    borde = Border(bottom=Side(style="medium", color=AZUL))
    for i, t in enumerate(titulos, 1):
        c = ws.cell(row=fila, column=i, value=t)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=AZUL)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = borde
    ws.row_dimensions[fila].height = 32


def _anchos(ws, anchos):
    for i, w in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _fila(ws, r, valores):
    """valores: lista de (valor, formato|None)."""
    for j, (v, fmt) in enumerate(valores, 1):
        c = ws.cell(row=r, column=j, value=round(v, 2) if isinstance(v, float) else v)
        if fmt:
            c.number_format = fmt


def excel(rep, desde, hasta, agrupar, negocio, num_orden):
    wb = Workbook()
    ws = wb.active
    ws.title = "Consolidado"
    ws["A1"] = f"{negocio} — Reporte {AGRUPACIONES[agrupar].lower()}"
    ws["A1"].font = Font(bold=True, size=14, color=AZUL)
    ws["A2"] = f"Del {desde.strftime('%d/%m/%Y')} al {hasta.strftime('%d/%m/%Y')}"
    ws["A2"].font = Font(color="5B6B80")
    titulos = ["Periodo"] + [t for _, t, _ in COLUMNAS] + ["Margen neto"]
    _encabezado(ws, 4, titulos)
    r = 5
    for f in rep["filas"]:
        ws.cell(row=r, column=1, value=f["periodo"])
        for j, (c, _, dinero) in enumerate(COLUMNAS, 2):
            cell = ws.cell(row=r, column=j, value=round(f[c], 2))
            cell.number_format = FMT_COP if dinero else FMT_NUM
        m = ws.cell(row=r, column=len(COLUMNAS) + 2, value=round(f["margen"], 1) if f["margen"] is not None else None)
        m.number_format = FMT_PCT
        if r % 2 == 0:
            for j in range(1, len(titulos) + 1):
                ws.cell(row=r, column=j).fill = PatternFill("solid", fgColor="F5F9FD")
        r += 1
    ultima = r - 1
    ws.cell(row=r, column=1, value="TOTAL").font = Font(bold=True)
    claves = [c for c, _, _ in COLUMNAS]
    for j, (c, _, dinero) in enumerate(COLUMNAS, 2):
        col = get_column_letter(j)
        cell = ws.cell(row=r, column=j, value=f"=SUM({col}5:{col}{ultima})")
        cell.number_format = FMT_COP if dinero else FMT_NUM
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=CELESTE)
    col_s = get_column_letter(claves.index("servicios") + 2)
    col_e = get_column_letter(claves.index("ventas_eq") + 2)
    col_n = get_column_letter(claves.index("neta") + 2)
    m = ws.cell(row=r, column=len(COLUMNAS) + 2,
                value=f'=IF(({col_s}{r}+{col_e}{r})=0,"",{col_n}{r}/({col_s}{r}+{col_e}{r})*100)')
    m.number_format = FMT_PCT
    m.font = Font(bold=True)
    ws.cell(row=r, column=1).fill = PatternFill("solid", fgColor=CELESTE)
    ws.cell(row=r, column=len(COLUMNAS) + 2).fill = PatternFill("solid", fgColor=CELESTE)
    ws.freeze_panes = "B5"
    _anchos(ws, [24] + [15] * len(COLUMNAS) + [12])
    if len(rep["filas"]) > 1:
        ch = BarChart()
        ch.title = "Ingresos por servicios y ganancia neta"
        ch.y_axis.numFmt = '"$" #,##0'
        ch.add_data(Reference(ws, min_col=claves.index("servicios") + 2, min_row=4, max_row=ultima), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=5, max_row=ultima))
        ln = LineChart()
        ln.add_data(Reference(ws, min_col=claves.index("neta") + 2, min_row=4, max_row=ultima), titles_from_data=True)
        ch += ln
        ch.height, ch.width = 9, 22
        ws.add_chart(ch, f"B{r + 3}")

    w = wb.create_sheet("Por servicio")
    _encabezado(w, 1, ["Tipo de servicio", "Órdenes", "Equipos atendidos", "Ingresos", "Ganancia"])
    for i, t in enumerate(rep["por_tipo"], 2):
        _fila(w, i, [(t["nombre"], None), (int(t["n"]), FMT_NUM), (int(t["equipos"]), FMT_NUM),
                     (t["ingresos"], FMT_COP), (t["ganancia"], FMT_COP)])
    _anchos(w, [32, 12, 18, 16, 16])

    w = wb.create_sheet("Por técnico")
    _encabezado(w, 1, ["Técnico", "Órdenes", "Visitas en garantía", "Ingresos", "Comisiones", "Ganancia"])
    for i, t in enumerate(rep["por_tecnico"], 2):
        _fila(w, i, [(t["nombre"], None), (int(t["n"]), FMT_NUM), (int(t.get("garantias", 0)), FMT_NUM),
                     (t["ingresos"], FMT_COP), (t["comision"], FMT_COP), (t["ganancia"], FMT_COP)])
    _anchos(w, [28, 12, 18, 16, 16, 16])

    w = wb.create_sheet("Por cliente")
    _encabezado(w, 1, ["Cliente", "Servicios", "Equipos comprados", "Ingresos"])
    for i, t in enumerate(rep["por_cliente"], 2):
        _fila(w, i, [(t["nombre"], None), (int(t.get("servicios", 0)), FMT_NUM),
                     (int(t.get("equipos", 0)), FMT_NUM), (t["ingresos"], FMT_COP)])
    _anchos(w, [34, 12, 18, 16])

    ini = datetime.combine(desde, datetime.min.time())
    fin = datetime.combine(hasta + timedelta(days=1), datetime.min.time())
    w = wb.create_sheet("Órdenes")
    _encabezado(w, 1, ["Terminada", "Orden", "Cliente", "Servicio", "Técnico", "Equipos", "Garantía",
                       "Cobrado", "Materiales", "Comisión", "Otros costos", "Ganancia", "Pago"])
    for i, o in enumerate(ordenes_periodo(ini, fin), 2):
        _fila(w, i, [(o.fecha_fin.date(), FMT_F), (num_orden(o), None), (o.cliente_nombre or "", None),
                     (o.tipo_txt, None), (o.tecnico.nombre if o.tecnico else "", None), (len(o.equipos), FMT_NUM),
                     ("Sí" if o.en_garantia else "", None), (o.total, FMT_COP), (o.costo_materiales, FMT_COP),
                     (o.pago_tecnico, FMT_COP), (o.otros_costos, FMT_COP), (o.ganancia, FMT_COP),
                     ("Crédito" if o.condicion == "credito" else (o.metodo_pago or ""), None)])
    _anchos(w, [12, 12, 28, 26, 20, 9, 9, 14, 14, 14, 14, 14, 14])

    w = wb.create_sheet("Venta de equipos")
    _encabezado(w, 1, ["Fecha", "N.º", "Cliente", "Equipo", "Cant.", "Proveedor", "Costo", "Otros costos",
                       "Precio venta", "Ganancia", "Margen"])
    for i, v in enumerate(ventas_eq_periodo(ini, fin), 2):
        _fila(w, i, [(v.fecha.date(), FMT_F), (v.numero, None), (v.cliente_nombre or "", None),
                     (v.descripcion, None), (v.cantidad, FMT_NUM), (v.proveedor or "", None),
                     (v.costo_total, FMT_COP), (v.otros_costos, FMT_COP), (v.precio_total, FMT_COP),
                     (v.ganancia, FMT_COP), (f'=IF(I{i}=0,"",J{i}/I{i}*100)', FMT_PCT)])
    _anchos(w, [12, 11, 26, 38, 7, 22, 14, 14, 14, 14, 10])

    w = wb.create_sheet("Gastos")
    _encabezado(w, 1, ["Categoría", "Total"])
    for i, (cat, v) in enumerate(rep["gastos_cat"], 2):
        _fila(w, i, [(cat, None), (v, FMT_COP)])
    base = len(rep["gastos_cat"]) + 4
    _encabezado(w, base, ["Fecha", "Categoría", "Descripción", "Valor"])
    for i, g in enumerate(Gasto.query.filter(Gasto.fecha >= desde, Gasto.fecha <= hasta)
                          .order_by(Gasto.fecha).all(), base + 1):
        _fila(w, i, [(g.fecha, FMT_F), (g.categoria, None), (g.descripcion or "", None), (g.valor, FMT_COP)])
    _anchos(w, [24, 26, 40, 16])

    w = wb.create_sheet("Compras")
    _encabezado(w, 1, ["Fecha", "Material / herramienta", "Cantidad", "Unidad", "Costo unitario", "Total", "Proveedor"])
    i = 2
    for c in (CompraInsumo.query.filter(CompraInsumo.anulada.is_(False), CompraInsumo.fecha >= ini,
                                        CompraInsumo.fecha < fin).order_by(CompraInsumo.fecha).all()):
        _fila(w, i, [(c.fecha.date(), FMT_F), (c.insumo.nombre, None), (c.cantidad, "#,##0.###"),
                     (c.insumo.unidad, None), (c.costo_unitario, FMT_COP), (c.costo_total, FMT_COP),
                     (c.proveedor or "", None)])
        i += 1
    for a in (Activo.query.filter(Activo.eliminado.is_(False), Activo.fecha_compra >= desde,
                                  Activo.fecha_compra <= hasta).order_by(Activo.fecha_compra).all()):
        _fila(w, i, [(a.fecha_compra, FMT_F), (f"Herramienta: {a.nombre}", None), (1, FMT_NUM), ("unidad", None),
                     (a.valor, FMT_COP), (a.valor, FMT_COP), (a.proveedor or "", None)])
        i += 1
    _anchos(w, [12, 34, 12, 12, 15, 15, 26])

    w = wb.create_sheet("Cartera")
    _encabezado(w, 1, ["Documento", "Fecha", "Cliente", "Vence", "Total", "Abonado", "Saldo", "Estado"])
    docs = [o for o in Orden.query.filter(Orden.anulada.is_(False), Orden.condicion == "credito",
                                          Orden.estado == "terminada").all() if o.saldo > 0.5]
    docs += [v for v in VentaEquipo.query.filter(VentaEquipo.anulada.is_(False),
                                                 VentaEquipo.condicion == "credito").all() if v.saldo > 0.5]
    docs.sort(key=lambda x: x.vence or date.max)
    for i, x in enumerate(docs, 2):
        es_orden = isinstance(x, Orden)
        _fila(w, i, [(num_orden(x) if es_orden else x.numero, None),
                     ((x.fecha_fin if es_orden else x.fecha).date(), FMT_F), (x.cliente_nombre or "", None),
                     (x.vence, FMT_F), (x.total_factura, FMT_COP), (x.pagado, FMT_COP), (x.saldo, FMT_COP),
                     (x.estado_pago[0], None)])
    _anchos(w, [14, 12, 28, 12, 14, 14, 14, 18])

    w = wb.create_sheet("Agenda preventivos")
    _encabezado(w, 1, ["Próximo preventivo", "Cliente", "Teléfono", "Dirección", "Equipo", "Último", "Frecuencia (meses)"])
    hoy = ahora().date()
    eqs = (Equipo.query.filter(Equipo.activo.is_(True), Equipo.frecuencia_meses > 0,
                               Equipo.proximo_preventivo.isnot(None),
                               Equipo.proximo_preventivo <= hoy + timedelta(days=90))
           .order_by(Equipo.proximo_preventivo).all())
    for i, e in enumerate(eqs, 2):
        _fila(w, i, [(e.proximo_preventivo, FMT_F), (e.cliente.nombre, None), (e.cliente.telefono or "", None),
                     (e.cliente.direccion_completa, None), (e.nombre, None), (e.ultimo_preventivo, FMT_F),
                     (e.frecuencia_meses, FMT_NUM)])
    _anchos(w, [16, 28, 16, 34, 40, 12, 12])

    for hoja in wb.worksheets[1:]:
        hoja.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
